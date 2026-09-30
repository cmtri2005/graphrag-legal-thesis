"""Labels outside the D4 cross-check: closed vocabulary, blind cross-check rows."""
import csv
import json

import pytest

from labels import COLUMNS, add, from_raw, read


def _data(tmp_path, crosscheck_verdicts=("", "")):
    (tmp_path / "derived/vilextime").mkdir(parents=True)
    (tmp_path / "review").mkdir()
    rows = [{"id": f"vilextime:T2:p{i}:2020-01-01", "group": "T2"} for i in range(3)]
    (tmp_path / "derived/vilextime/release.jsonl").write_text("".join(json.dumps(r) + "\n" for r in rows))
    for name, verdict in zip(("crosscheck_A.tsv", "crosscheck_B.tsv"), crosscheck_verdicts):
        (tmp_path / "review" / name).write_text(f"id\tverdict\n{rows[0]['id']}\t{verdict}\n")
    return tmp_path, [r["id"] for r in rows]


def _batch(path, labels):
    with path.open("w", newline="") as f:
        w = csv.DictWriter(f, COLUMNS, delimiter="\t", extrasaction="ignore")
        w.writeheader()
        w.writerows(labels)
    return path


def test_model_label_is_refused_on_a_crosscheck_row_until_both_people_labelled(tmp_path):
    data, ids = _data(tmp_path)
    batch = _batch(tmp_path / "b.tsv", [{"id": ids[0], "labeler": "claude-opus-5-5", "verdict": "dat"}])
    with pytest.raises(SystemExit, match="kiểm chéo"):
        add(data, batch)
    data, ids = _data(tmp_path / "done", ("dat", "khong_dat"))
    add(data, batch)
    assert [r["id"] for r in read(data / "review/labels.tsv")] == [ids[0]]


def test_out_of_vocabulary_rejects_the_whole_batch(tmp_path):
    data, ids = _data(tmp_path)
    batch = _batch(tmp_path / "b.tsv", [
        {"id": ids[1], "labeler": "claude-opus-5-5", "verdict": "dat"},
        {"id": ids[2], "labeler": "claude-opus-5-5", "verdict": "khong_dat", "reasons": "looks_wrong"},
    ])
    with pytest.raises(SystemExit, match="reasons ngoài từ vựng"):
        add(data, batch)
    assert not (data / "review/labels.tsv").exists()


def test_every_model_khong_dat_is_audited(tmp_path):
    data, ids = _data(tmp_path)
    add(data, _batch(tmp_path / "b.tsv", [
        {"id": ids[1], "labeler": "claude-opus-5-5", "verdict": "khong_dat", "reasons": "question_ambiguous"},
        {"id": ids[2], "labeler": "human:A", "verdict": "khong_dat", "reasons": "gold_wrong"},
    ]))
    labels = {r["labeler"]: r for r in read(data / "review/labels.tsv")}
    assert labels["claude-opus-5-5"]["audit_required"] == "có"
    assert labels["human:A"]["audit_required"] == ""


def test_workbook_cells_map_onto_the_vocabulary():
    assert from_raw("verified \n(đã chỉnh tổ chức -> cá nhân)") == ("dat", [], [], "đã chỉnh tổ chức -> cá nhân")
    assert from_raw("suspected: question-answer") == ("khong_dat", ["unanswerable_as_asked"], [], "")
    assert from_raw("suspected: valid_to + group + num_versions") == ("", [], ["group_wrong", "valid_to_wrong"], "")
    assert from_raw("suspected") == ("", [], [], "")


def test_audit_round_trips_through_the_workbook(tmp_path):
    from openpyxl import load_workbook

    from labels import audit_sheet, import_audit

    data, ids = _data(tmp_path)
    rows = [json.loads(line) for line in (data / "derived/vilextime/release.jsonl").open()]
    for r in rows:
        r.update(question="q", as_of="2020-01-01", answer="a", source_url="u", target={"citation": "c", "law_id": "l"})
    (data / "derived/vilextime/release.jsonl").write_text("".join(json.dumps(r) + "\n" for r in rows))
    add(data, _batch(tmp_path / "b.tsv", [
        {"id": ids[1], "labeler": "claude-opus-5-5", "verdict": "khong_dat", "reasons": "question_ambiguous"}]))
    audit_sheet(data, tmp_path / "a.xlsx", None)
    wb = load_workbook(tmp_path / "a.xlsx")
    ws = wb.active
    head = [c.value for c in ws[1]]
    ws.cell(2, head.index("audit") + 1, "khong_dong_y")
    ws.cell(2, head.index("auditor") + 1, "claude-opus-5-5")
    wb.save(tmp_path / "a.xlsx")
    with pytest.raises(SystemExit):                     # a model never audits (guideline 1.2 §6)
        import_audit(data, tmp_path / "a.xlsx")
    ws.cell(2, head.index("auditor") + 1, "human:A")
    wb.save(tmp_path / "a.xlsx")
    import_audit(data, tmp_path / "a.xlsx")
    (label,) = read(data / "review/labels.tsv")
    assert (label["audit"], label["auditor"]) == ("khong_dong_y", "human:A")


def test_a_row_whose_lineage_was_dropped_still_reaches_the_audit_sheet(tmp_path):
    from openpyxl import load_workbook

    from labels import audit_sheet

    data, ids = _data(tmp_path)
    add(data, _batch(tmp_path / "b.tsv", [
        {"id": ids[1], "labeler": "claude-opus-5-5", "verdict": "khong_dat", "reasons": "group_mismatch"}]))
    gone = {"id": ids[1], "group": "T2", "lineage_id": "p1", "question": "q", "as_of": "2020-01-01", "answer": "a",
            "source_url": "u", "target": {"citation": "c", "law_id": "l"}}
    (data / "derived/vilextime/T2.jsonl").write_text(json.dumps(gone) + "\n")
    kept = [line for line in (data / "derived/vilextime/release.jsonl").open() if ids[1] not in line]
    (data / "derived/vilextime/release.jsonl").write_text("".join(kept))
    (data / "vilextime").mkdir()
    (data / "vilextime/dropped.jsonl").write_text(json.dumps({"lineage_id": "p1", "reason": "group_mismatch",
                                                              "note": "hai đáp án"}) + "\n")
    audit_sheet(data, tmp_path / "a.xlsx", None)
    ws = load_workbook(tmp_path / "a.xlsx").active
    head, row = [c.value for c in ws[1]], [c.value for c in ws[2]]
    assert row[head.index("id")] == ids[1]
    assert row[head.index("fix")] == "loại lineage (group_mismatch): hai đáp án"


def test_guideline_1_2_audits_transitional_and_district_rows_and_its_own_labels():
    from labels import audit_required

    def model(i, gv="1.1"):
        return {"id": f"vilextime:T1:provision:x{i}:2020-01-02", "labeler": "claude-opus-5-5", "verdict": "dat",
                "meta_issues": "", "guideline_version": gv}
    plain = next(model(i) for i in range(50) if not audit_required(model(i)))   # outside the 20% sample
    assert audit_required(plain, {"legal_markers": ["actor_has_transitional"]}) == "có"
    assert audit_required(plain, {"as_of": "2025-08-01", "question": "UBND cấp huyện cấp giấy?"}) == "có"
    assert audit_required(plain, {"as_of": "2024-08-01", "question": "UBND cấp huyện cấp giấy?"}) == ""
    assert audit_required({**plain, "guideline_version": "1.2"}) == "có"

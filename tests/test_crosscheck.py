"""Agreement on the D4 cross-check: AC1 as the gate, Cohen κ reported beside it."""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts" / "review"))
from crosscheck import (  # noqa: E402
    AC1_TARGET, PNEG_TARGET, agreement_needed, bootstrap_ci, cohen_kappa, gwet_ac1, pair_scores,
    specific_agreement,
)


def test_kappa_matches_a_hand_computed_case():
    # 10 items, 7 agree -> po = 0.7; A has 6 dat, B has 7 -> pe = 0.6*0.7 + 0.4*0.3 = 0.54
    a = ["dat"] * 5 + ["khong_dat"] * 2 + ["dat", "khong_dat", "khong_dat"]
    b = ["dat"] * 5 + ["khong_dat"] * 2 + ["khong_dat", "dat", "dat"]
    assert abs(cohen_kappa(a, b) - (0.7 - 0.54) / (1 - 0.54)) < 1e-9


def test_kappa_edges():
    assert cohen_kappa(["dat", "khong_dat"], ["dat", "khong_dat"]) == 1.0
    assert cohen_kappa(["dat", "dat"], ["dat", "dat"]) is None       # all chance: undefined
    assert cohen_kappa([], []) is None


def test_ac1_stays_high_when_both_raters_say_dat_almost_always():
    # 20 items, both say dat on 18, disagree on 2: raw agreement 0.9, κ collapses, AC1 does not.
    a = ["dat"] * 19 + ["khong_dat"]
    b = ["dat"] * 18 + ["khong_dat", "dat"]
    assert cohen_kappa(a, b) < 0                       # the κ paradox on a skewed sample
    assert gwet_ac1(a, b) > 0.85
    assert gwet_ac1(["dat", "khong_dat"], ["dat", "khong_dat"]) == 1.0


def test_p_neg_reports_agreement_on_the_verdict_that_drops_a_row():
    """`dat` agreement is nearly free at this prevalence; `khong_dat` is the claim."""
    a = ["dat"] * 17 + ["khong_dat"] * 3
    b = ["dat"] * 18 + ["khong_dat"] * 2          # they overlap on 2 of the khong_dat rows

    got = specific_agreement(a, b)
    assert got["dat"] > 0.9                        # the easy half
    assert got["khong_dat"] == 2 * 2 / (3 + 2)     # 2 * both / (A's + B's)
    assert specific_agreement(["dat"], ["khong_dat"]) == {"dat": 0.0, "khong_dat": 0.0}


def test_kappa_carries_the_agreement_its_target_would_need():
    """The prevalence paradox, as a number rather than as an argument.

    At the D4 marginals κ ≥ 0.60 needs p_o ≥ 0.843, and raters who both settle
    near 85% `dat` need p_o ≥ 0.898 — at most ~30 disagreements in 300. No extra
    calibration moves that, which is why AC1 is the pre-registered gate.
    """
    skewed = (["dat"] * 85 + ["khong_dat"] * 15, ["dat"] * 85 + ["khong_dat"] * 15)
    needed = agreement_needed(*skewed, 0.60)
    assert 0.89 < needed < 0.91

    balanced = agreement_needed(["dat"] * 50 + ["khong_dat"] * 50,
                                ["dat"] * 50 + ["khong_dat"] * 50, 0.60)
    assert needed > balanced, "biên lệch phải đòi p_o cao hơn biên cân"
    assert agreement_needed([], [], 0.60) is None


def test_the_bootstrap_interval_is_seeded_and_brackets_the_point_estimate():
    a = ["dat"] * 40 + ["khong_dat"] * 10
    b = ["dat"] * 38 + ["khong_dat"] * 12

    low, high = bootstrap_ci(a, b, gwet_ac1, draws=200)
    assert low <= gwet_ac1(a, b) <= high
    assert bootstrap_ci(a, b, gwet_ac1, draws=200) == [low, high]   # seeded: reruns match
    assert bootstrap_ci(["dat"], ["dat"], gwet_ac1) is None         # one item: no interval


def test_perfect_agreement_passes_the_gate_and_a_skewed_split_does_not():
    perfect = ["dat"] * 40 + ["khong_dat"] * 10
    scored = pair_scores(perfect, list(perfect))
    assert scored["ac1"] >= AC1_TARGET and scored["p_neg"] >= PNEG_TARGET

    # agree on every `dat`, never on `khong_dat`: p_o 0.80 yet the gate must fail
    missed = pair_scores(["dat"] * 40 + ["khong_dat"] * 10, ["dat"] * 50)
    assert missed["agreement"] == 0.8 and missed["p_neg"] == 0.0
    assert missed["p_neg"] < PNEG_TARGET


def _release_row(lid, group, as_of, question="Hỏi?", **extra):
    return {"id": f"vilextime:{group}:{lid}:{as_of}", "group": group, "lineage_id": lid, "question": question,
            "as_of": as_of, "as_of_strategy": "ranh_gioi", "answer": "Đáp.", "no_answer_reason": None,
            "changed_span": "", "valid_from": "2020-01-01", "valid_to": None, "source_url": "u",
            "target": {"lead_in": "", "citation": "Điều 1", "law_id": "1/2020/TT"}, **extra}


def test_t6_takes_what_t5_leaves_of_150_in_the_release_proportion():
    from crosscheck import targets

    rows = {"t5": [{"group": "T5"}] * 38,
            **{f"a{i}": [{"group": "T6", "repeal_mode": "ancestor"}] * 2 for i in range(150)},
            **{f"d{i}": [{"group": "T6", "repeal_mode": "direct"}] * 2 for i in range(44)}}
    t = targets(rows)
    assert (t[("T6", "ancestor")], t[("T6", "direct")]) == (86, 26) and t[("T5", None)] is None


def test_build_keeps_a_label_only_with_its_question_and_answer_and_saves_the_old_files(tmp_path):
    import csv
    import json

    import pytest
    from crosscheck import COLUMNS, build

    (tmp_path / "derived/vilextime").mkdir(parents=True)
    (tmp_path / "review").mkdir()
    rows = [_release_row("p", "T2", "2021-01-01"), _release_row("p", "T2", "2021-01-02"),
            _release_row("q", "T3", "2021-01-01", question="Mới?")]
    (tmp_path / "derived/vilextime/release.jsonl").write_text(
        "".join(json.dumps(r, ensure_ascii=False) + "\n" for r in rows), encoding="utf-8")
    for name, extra in (("crosscheck_A.tsv", ["adjudicated"]), ("crosscheck_B.tsv", [])):
        with (tmp_path / "review" / name).open("w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, COLUMNS + extra, delimiter="\t", lineterminator="\n", restval="")
            w.writeheader()
            # p was T4 when it was labelled: the id changed, the lineage and date did not
            w.writerow({"id": "vilextime:T4:p:2021-01-01", "lineage_id": "p", "as_of": "2021-01-01",
                        "question": "Hỏi?", "answer": "Đáp.", "verdict": "dat"})
            w.writerow({"id": rows[2]["id"], "lineage_id": "q", "as_of": "2021-01-01", "question": "Cũ?",
                        "answer": "Đáp.", "verdict": "khong_dat"})
    build(tmp_path, today="20261011")
    for name in ("crosscheck_A", "crosscheck_B"):
        assert (tmp_path / f"review/{name}.20261011.tsv").exists()
        got = {r["id"]: r["verdict"] for r in csv.DictReader((tmp_path / f"review/{name}.tsv").open(), delimiter="\t")}
        assert got == {rows[0]["id"]: "dat", rows[1]["id"]: "", rows[2]["id"]: ""}   # q's question changed
        version = {r["id"]: r["guideline_version"] for r in csv.DictReader((tmp_path / f"review/{name}.tsv").open(),
                                                                              delimiter="\t")}
        assert version[rows[0]["id"]] == "1.2"                                       # labelled before versions
    with pytest.raises(SystemExit):
        build(tmp_path, today="20261011")                                             # never overwrite a saved copy


def test_export_refuses_only_when_the_workbook_holds_labels_not_yet_imported(tmp_path):
    import csv

    import pytest
    from openpyxl import load_workbook

    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts" / "review"))
    from crosscheck import COLUMNS
    from crosscheck_xlsx import SHEET, export

    (tmp_path / "review").mkdir()
    for who, extra in (("A", ["adjudicated"]), ("B", [])):
        with (tmp_path / f"review/crosscheck_{who}.tsv").open("w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, COLUMNS + extra, delimiter="\t", lineterminator="\n", restval="")
            w.writeheader()
            w.writerow({"id": "r1", "lineage_id": "p", "question": "Hỏi?", "verdict": "dat"})   # carried over
            w.writerow({"id": "r2", "lineage_id": "q", "question": "Hỏi?"})
    export(tmp_path)                                   # a label in the TSV alone does not block
    path = tmp_path / "review/crosscheck_A.xlsx"
    wb = load_workbook(path)
    ws = wb[SHEET]
    col = [c.value for c in ws[1]].index("verdict") + 1
    rows = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}
    assert ws.cell(rows["r1"], col).protection.locked and not ws.cell(rows["r2"], col).protection.locked
    ws.cell(rows["r2"], col, "khong_dat")             # labelled in the workbook, not imported yet
    wb.save(path)
    with pytest.raises(SystemExit):
        export(tmp_path)


def test_import_stamps_todays_guideline_only_on_labels_made_now(tmp_path):
    import csv

    from openpyxl import load_workbook

    from crosscheck import COLUMNS, GUIDELINE, read
    from crosscheck_xlsx import SHEET, export, ingest

    (tmp_path / "review").mkdir()
    for who, extra in (("A", ["adjudicated"]), ("B", [])):
        with (tmp_path / f"review/crosscheck_{who}.tsv").open("w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, COLUMNS + extra, delimiter="\t", lineterminator="\n", restval="")
            w.writeheader()
            w.writerow({"id": "r1", "lineage_id": "p", "question": "Hỏi?", "verdict": "dat"})   # a 1.2 label
            w.writerow({"id": "r2", "lineage_id": "q", "question": "Hỏi?"})
    export(tmp_path)
    path = tmp_path / "review/crosscheck_A.xlsx"
    wb = load_workbook(path)
    ws = wb[SHEET]
    col = [c.value for c in ws[1]].index("verdict") + 1
    ws.cell({ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}["r2"], col, "dat")
    wb.save(path)
    ingest(tmp_path, "A")
    got = {r["id"]: r["guideline_version"] for r in read(tmp_path / "review/crosscheck_A.tsv")}
    assert got == {"r1": "1.2", "r2": GUIDELINE}


def test_pilot_takes_whole_unlabelled_t6_lineages_outside_d4_and_never_overwrites(tmp_path):
    import csv
    import json

    import pytest
    from crosscheck import COLUMNS, PILOT_ROWS, pilot, read

    (tmp_path / "derived/vilextime").mkdir(parents=True)
    (tmp_path / "review").mkdir()
    rows = [_release_row(f"{mode[0]}{i}", "T6", day, repeal_mode=mode)
            for mode, n in (("ancestor", 30), ("direct", 10)) for i in range(n) for day in ("2021-01-01", "2021-06-01")]
    (tmp_path / "derived/vilextime/release.jsonl").write_text(
        "".join(json.dumps(r, ensure_ascii=False) + "\n" for r in rows), encoding="utf-8")
    with (tmp_path / "review/crosscheck_A.tsv").open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, COLUMNS, delimiter="\t", lineterminator="\n", restval="")
        w.writeheader()
        w.writerow({"id": rows[0]["id"]})                                  # a0 is in D4
    (tmp_path / "review/labels.tsv").write_text(f"id\tlabeler\n{rows[3]['id']}\tclaude\n", encoding="utf-8")   # a1
    pilot(tmp_path)
    got = read(tmp_path / "review/pilot_T6_A.tsv")
    lineages = {r["lineage_id"] for r in got}
    assert len(got) == PILOT_ROWS and not {"a0", "a1"} & lineages
    assert sum(lid.startswith("a") for lid in lineages) == 7                # 28 of 38 free pairs -> 14.7 -> 14 rows
    assert got == read(tmp_path / "review/pilot_T6_B.tsv")
    with pytest.raises(SystemExit):
        pilot(tmp_path)


def test_kiem_covers_every_boundary_row_overlaps_a_quarter_and_carries_d4_only_unchanged(tmp_path):
    import csv
    import json

    from crosscheck import COLUMNS, kiem, read

    (tmp_path / "derived/vilextime").mkdir(parents=True)
    (tmp_path / "review").mkdir()
    rows = [_release_row(f"p{i}", "T2", day, version_id=f"version:p{i}:{k}")
            for i in range(8) for k, day in enumerate(("2021-01-01", "2021-06-01"), 1)]
    rows.append({**_release_row("p0", "T2", "2021-03-01", version_id="version:p0:1"),
                 "as_of_strategy": "giua_khoang"})                                       # a twin: judged with its row
    rows[2]["short_answer"] = "Đáp"                    # p1's two versions read alike: no contrast to score
    (tmp_path / "derived/vilextime/release.jsonl").write_text(
        "".join(json.dumps(r, ensure_ascii=False) + "\n" for r in rows), encoding="utf-8")
    for name, extra in (("crosscheck_A.tsv", ["adjudicated"]), ("crosscheck_B.tsv", [])):
        with (tmp_path / "review" / name).open("w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, COLUMNS + extra, delimiter="\t", lineterminator="\n", restval="")
            w.writeheader()
            for day in ("2021-01-01", "2021-06-01"):    # D4 labelled p0 alike on both sides
                w.writerow({"id": f"vilextime:T2:p0:{day}", "lineage_id": "p0", "as_of": day,
                            "question": "Hỏi?", "answer": "Đáp.", "verdict": "dat", "guideline_version": "2.0"})
            w.writerow({"id": "vilextime:T2:p1:2021-01-01", "lineage_id": "p1", "as_of": "2021-01-01",
                        "question": "Cũ?", "answer": "Đáp.", "verdict": "dat"})            # question changed since
    kiem(tmp_path, today="20261020")
    a, b = read(tmp_path / "review/kiem_A.tsv"), read(tmp_path / "review/kiem_B.tsv")
    ids_a, ids_b = {r["id"] for r in a}, {r["id"] for r in b}
    assert ids_a | ids_b == {r["id"] for r in rows if r["as_of_strategy"] == "ranh_gioi"}
    lo_a, lo_b = {r["id"]: r["lo"] for r in a}, {r["id"]: r["lo"] for r in b}
    calib = {i for i, lo in lo_a.items() if lo == "0"}
    assert calib == {i for i, lo in lo_b.items() if lo == "0"} and len(calib) == 4   # CALIB: 2 T2 regular lineages
    assert len((ids_a & ids_b) - calib) == 4                        # 25% of 16 rows, whole lineages
    assert all(lo_a[i] == lo_b[i] for i in ids_a & ids_b)           # a shared row sits in the same batch
    assert not {"vilextime:T2:p0:2021-01-01"} & ids_a & ids_b       # D4 cannot measure agreement again
    labelled = {r["id"]: (r["verdict"], r["guideline_version"]) for r in a + b if r["verdict"]}
    assert labelled == {f"vilextime:T2:p0:{d}": ("dat", "2.0") for d in ("2021-01-01", "2021-06-01")}
    assert {r["lo"] for r in a} == {"0", "1"}
    hints = {r["id"]: r["note_for_annotator"] for r in a + b}
    assert "đáp án ngắn cũng nằm trong phiên bản 2" in hints[rows[2]["id"]] and not hints[rows[3]["id"]]
    kiem(tmp_path, today="20261021")                               # a re-run keeps who has which lineage
    assert {r["id"]: r["lo"] for r in read(tmp_path / "review/kiem_A.tsv")} == lo_a


def test_batches_spread_the_overlap_evenly_and_fill_with_singles():
    from crosscheck import batches

    size = {**{f"o{i}": 2 for i in range(6)}, **{f"s{i}": 1 for i in range(18)}}
    lots = batches([f"o{i}" for i in range(6)], [f"s{i}" for i in range(18)], size, 10)
    assert len(lots) == 3 and [sum(lid.startswith("o") for lid in lot) for lot in lots] == [2, 2, 2]
    assert sorted(lid for lot in lots for lid in lot) == sorted(size)
    assert [sum(size[lid] for lid in lot) for lot in lots] == [10, 10, 10]
    uneven = batches(["o0", "o1"], [f"s{i}" for i in range(17)], size, 10)   # 21 rows: 3 batches of 7
    assert [sum(size[lid] for lid in lot) for lot in uneven] == [7, 7, 7]


def test_kappa_leaves_the_calibration_batch_out_of_agreement(tmp_path):
    import csv
    import json

    from crosscheck import GUIDELINE, KIEM_COLUMNS, kappa

    (tmp_path / "review").mkdir()
    for who, verdicts in (("A", ("dat", "dat", "dat", "dat")), ("B", ("khong_dat", "dat", "dat", "khong_dat"))):
        with (tmp_path / f"review/kiem_{who}.tsv").open("w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, KIEM_COLUMNS, delimiter="\t", lineterminator="\n", restval="")
            w.writeheader()
            for k, (lo, verdict) in enumerate(zip(("0", "1", "1", "1"), verdicts)):
                w.writerow({"lo": lo, "id": f"r{k}", "group": "T2", "lineage_id": f"p{k}", "verdict": verdict,
                            "reasons": "group_mismatch" if verdict == "khong_dat" else "",
                            "guideline_version": "2.0" if k == 3 else GUIDELINE})   # r3: a D4 label, not counted
    kappa(tmp_path, "kiem")
    result = json.loads((tmp_path / "review/kiem_result.json").read_text(encoding="utf-8"))
    assert result["labelled_by_both"] == 2 and result["disagreements"] == []
    assert result["calibration"] == {"n": 1, "disagreements": ["r0"]}


def test_changed_words_are_highlighted_and_the_import_reads_the_plain_text_back(tmp_path):
    import csv

    from openpyxl import load_workbook
    from openpyxl.cell.rich_text import CellRichText

    from crosscheck import KIEM_COLUMNS, read
    from crosscheck_xlsx import SHEET, changed_spans, export, ingest

    text = "Thời hạn  là 15 ngày;\nkể từ ngày nhận."
    spans = changed_spans(text, ["Thời hạn là 10 ngày; kể từ ngày nhận."])
    assert [text[s:e] for s, e in spans] == ["15"]
    (tmp_path / "review").mkdir()
    for who in "AB":
        with (tmp_path / f"review/kiem_{who}.tsv").open("w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, KIEM_COLUMNS, delimiter="\t", lineterminator="\n", restval="")
            w.writeheader()
            for k, answer in enumerate((text, "Thời hạn là 10 ngày; kể từ ngày nhận.")):
                w.writerow({"lo": "1", "id": f"r{k}", "group": "T2", "lineage_id": "p", "question": "Hỏi?",
                            "answer": answer})
    export(tmp_path, "kiem")
    path = tmp_path / "review/kiem_A.xlsx"
    ws = load_workbook(path, rich_text=True)[SHEET]
    head = [c.value for c in ws[1]]
    cell = ws.cell(2, head.index("answer") + 1).value
    assert isinstance(cell, CellRichText) and str(cell) == text
    assert "Cách quyết" in load_workbook(path).sheetnames
    ingest(tmp_path, "A", "kiem")                                  # nothing labelled, nothing altered: no error
    assert [r["answer"] for r in read(tmp_path / "review/kiem_A.tsv")][0] == text


def test_kiem_import_refuses_a_new_label_without_the_required_fields(tmp_path):
    import csv

    import pytest
    from openpyxl import load_workbook

    from crosscheck import GUIDELINE, KIEM_COLUMNS, read
    from crosscheck_xlsx import SHEET, export, ingest

    (tmp_path / "review").mkdir()
    for who in "AB":
        with (tmp_path / f"review/kiem_{who}.tsv").open("w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, KIEM_COLUMNS, delimiter="\t", lineterminator="\n", restval="")
            w.writeheader()
            w.writerow({"lo": "1", "id": "r1", "group": "T6", "lineage_id": "p", "question": "Hỏi?",
                        "verdict": "dat", "note": "[D4] ", "guideline_version": "2.0"})   # carried from D4, left alone
            w.writerow({"lo": "1", "id": "r2", "group": "T6", "lineage_id": "q", "question": "Hỏi?"})
    export(tmp_path, "kiem")
    path = tmp_path / "review/kiem_A.xlsx"

    def fill(**cells):
        wb = load_workbook(path)
        ws = wb[SHEET]
        head = [c.value for c in ws[1]]
        row = {ws.cell(r, 2).value: r for r in range(2, ws.max_row + 1)}["r2"]
        for k, v in cells.items():
            ws.cell(row, head.index(k) + 1, v)
        wb.save(path)

    fill(verdict="dat", da_xem_van_ban_tac_dong="12/2020/NĐ-CP", phut="6")
    with pytest.raises(SystemExit):
        ingest(tmp_path, "A", "kiem")                              # T6 needs noi_dung_con_o_dau
    fill(noi_dung_con_o_dau="không có")
    ingest(tmp_path, "A", "kiem")
    got = {r["id"]: (r["verdict"], r["guideline_version"]) for r in read(tmp_path / "review/kiem_A.tsv")}
    assert got == {"r1": ("dat", "2.0"), "r2": ("dat", GUIDELINE)}
    fill(verdict=None, da_xem_van_ban_tac_dong=None, phut=None, noi_dung_con_o_dau=None)   # an xlsx older than the TSV
    with pytest.raises(SystemExit):
        ingest(tmp_path, "A", "kiem")
    assert read(tmp_path / "review/kiem_A.tsv")[1]["verdict"] == "dat"
    path = tmp_path / "review/kiem_B.xlsx"                         # B only measures agreement: verdict is enough
    fill(verdict="khong_dat", reasons="relocated", note="Điều 13 TT 57/2018")
    ingest(tmp_path, "B", "kiem")
    assert read(tmp_path / "review/kiem_B.tsv")[1]["verdict"] == "khong_dat"


def test_drop_adds_each_d4_lineage_with_a_final_khong_dat_once(tmp_path):
    import csv
    import json

    from adjudicate import drop
    from crosscheck import COLUMNS

    (tmp_path / "review").mkdir()
    (tmp_path / "vilextime").mkdir()
    for name, extra in (("crosscheck_A.tsv", ["adjudicated"]), ("crosscheck_B.tsv", [])):
        with (tmp_path / "review" / name).open("w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, COLUMNS + extra, delimiter="\t", lineterminator="\n", restval="")
            w.writeheader()
            w.writerow({"id": "r1", "group": "T2", "lineage_id": "p", "as_of": "2021-01-01",
                        "verdict": "khong_dat", "reasons": "gold_wrong"})                 # agreed
            w.writerow({"id": "r2", "group": "T2", "lineage_id": "q", "as_of": "2021-01-01",
                        "verdict": "dat" if name.endswith("A.tsv") else "khong_dat"})     # adjudicated below
            w.writerow({"id": "r3", "group": "T2", "lineage_id": "s", "as_of": "2021-01-01", "verdict": "dat"})
    (tmp_path / "review/hoa_giai_D4.tsv").write_text("id\tverdict\treasons\tnote\tadjudicator\nr2\tdat\t\t\thuman:C\n")
    drop(tmp_path, today="2026-10-20")
    drop(tmp_path, today="2026-10-21")
    got = [json.loads(x) for x in (tmp_path / "vilextime/dropped.jsonl").read_text().splitlines()]
    assert [(d["lineage_id"], d["reason"]) for d in got] == [("p", "gold_wrong")]


def test_adjudicating_kiem_writes_final_labels_to_a_and_keeps_raw_labels(tmp_path):
    import csv

    from openpyxl import load_workbook

    import adjudicate
    from crosscheck import GUIDELINE, KIEM_A_COLUMNS, KIEM_COLUMNS, read

    (tmp_path / "review").mkdir()
    for who, verdicts in (("A", ("dat", "dat", "dat")), ("B", ("khong_dat", "khong_dat", "dat"))):
        with (tmp_path / f"review/kiem_{who}.tsv").open("w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, KIEM_A_COLUMNS if who == "A" else KIEM_COLUMNS, delimiter="\t",
                               lineterminator="\n", restval="")
            w.writeheader()
            for k, (lo, verdict) in enumerate(zip(("0", "1", "1"), verdicts)):
                w.writerow({"lo": lo, "id": f"r{k}", "group": "T2", "lineage_id": f"p{k}", "verdict": verdict,
                            "reasons": "group_mismatch" if verdict == "khong_dat" else "",
                            "guideline_version": GUIDELINE})
    assert [r["nguyen_nhan"] for r in adjudicate.disagreements(tmp_path, "kiem")] == ["lô 0 hiệu chuẩn", "hiểu khác"]
    adjudicate.export(tmp_path, "kiem")
    path = tmp_path / "review/hoa_giai_kiem.xlsx"
    wb = load_workbook(path)
    ws = wb[adjudicate.SHEET]
    head = [c.value for c in ws[1]]
    for row in range(2, 4):
        ws.cell(row, head.index("chot_verdict") + 1, "dat")
    wb.save(path)
    adjudicate.ingest(tmp_path, "luật sư", "kiem")
    a = {r["id"]: r for r in read(tmp_path / "review/kiem_A.tsv")}
    assert a["r0"]["adjudicated"] == a["r1"]["adjudicated"] == "dat" and a["r2"]["adjudicated"] == ""
    assert a["r0"]["verdict"] == "dat"          # nhãn thô không bị đụng
    assert read(tmp_path / "review/hoa_giai_kiem.tsv")[0]["adjudicator"] == "luật sư"


def test_label_ui_save_validates_like_the_xlsx_import_and_keeps_the_other_columns(tmp_path):
    import csv

    import label_ui
    from crosscheck import GUIDELINE, KIEM_COLUMNS, read

    path = tmp_path / "kiem_A.tsv"
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, KIEM_COLUMNS, delimiter="\t", lineterminator="\n", restval="")
        w.writeheader()
        w.writerow({"lo": "1", "id": "r0", "group": "T2", "lineage_id": "p0", "answer": "a", "question": "q"})
    good = {"verdict": "dat", "da_xem_van_ban_tac_dong": "NĐ 1", "cham_lan_sua": "2025-01-01", "phut": "2"}
    for bad in ({**good, "verdict": ""}, {**good, "reasons": "gold_wrong"}, {**good, "phut": ""},
                {**good, "verdict": "khong_dat", "reasons": "gold_wrong"}):
        assert label_ui.save(path, "r0", bad)[1]
    assert read(path)[0]["verdict"] == ""                         # nothing written on error
    row, errors = label_ui.save(path, "r0", good)
    assert not errors and read(path)[0]["verdict"] == "dat"
    assert read(path)[0]["question"] == "q" and read(path)[0]["guideline_version"] == GUIDELINE


def test_label_ui_holds_back_the_gold_lineages_until_the_gold_sheet_is_filled(tmp_path):
    import label_ui

    (tmp_path / "gold_sample.tsv").write_text("khoa\tid\tlineage_id\tstratum\tannotators\n"
                                              "G001\tx\tp1\tT1\tA\nG002\ty\tp2\tT1\tB\nG003\tz\tp3\tT1\tAB\n",
                                              encoding="utf-8")
    sheet = tmp_path / "gold_A.tsv"
    sheet.write_text("khoa\thieu_luc\nG001\tcòn\nG003\t\n", encoding="utf-8")
    assert label_ui.held_back(tmp_path, "A") == {"p1", "p3"}
    sheet.write_text("khoa\thieu_luc\nG001\tcòn\nG003\tkhông\n", encoding="utf-8")
    assert label_ui.held_back(tmp_path, "A") == set()


def test_label_ui_edits_reach_the_whole_lineage_and_the_source_files(tmp_path):
    import csv
    import json

    import label_ui
    from crosscheck import KIEM_COLUMNS, read

    (tmp_path / "vilextime").mkdir()
    (tmp_path / "vilextime/questions.jsonl").write_text(json.dumps(
        {"lineage_id": "p", "question": "Hỏi cũ?", "model": "m"}, ensure_ascii=False) + "\n", encoding="utf-8")
    (tmp_path / "vilextime/short_answers.jsonl").write_text("".join(json.dumps(
        {"lineage_id": "p", "version_id": v, "short_answer": s, "by": "m"}, ensure_ascii=False) + "\n"
        for v, s in (("p:1", "7 triệu"), ("p:2", "9 triệu"))), encoding="utf-8")
    path = tmp_path / "kiem_A.tsv"
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, KIEM_COLUMNS, delimiter="\t", lineterminator="\n", restval="")
        w.writeheader()
        for rid, ans, short in (("r1", "Phạt 7 triệu đồng.", "7 triệu"), ("r2", "Phạt 9 triệu đồng.", "9 triệu")):
            w.writerow({"lo": "1", "id": rid, "group": "T1", "lineage_id": "p", "question": "Hỏi cũ?",
                        "answer": ans, "short_answer": short})
    version_of = {"r1": "p:1", "r2": "p:2"}
    label = {"verdict": "dat", "da_xem_van_ban_tac_dong": "không có", "phut": "1"}
    save = lambda edits: label_ui.save(path, "r1", label, edits, tmp_path, version_of)

    assert save({"short_answer": "10 triệu"})[1]                     # not in the wording: refused, nothing written
    assert read(path)[0]["verdict"] == "" and "10 triệu" not in (tmp_path / "vilextime/short_answers.jsonl").read_text()
    rows, errors = save({"question": "Hỏi mới?", "short_answer": "Phạt 7 triệu"})
    assert not errors and {r["id"] for r in rows} == {"r1", "r2"}
    got = {r["id"]: r for r in read(path)}
    assert got["r1"]["question"] == got["r2"]["question"] == "Hỏi mới?"
    assert (got["r1"]["short_answer"], got["r2"]["short_answer"]) == ("Phạt 7 triệu", "9 triệu")   # its version only
    q = json.loads((tmp_path / "vilextime/questions.jsonl").read_text(encoding="utf-8"))
    assert (q["question"], q["model_question"], q["edited_by"]) == ("Hỏi mới?", "Hỏi cũ?", "human:A")
    s = [json.loads(x) for x in (tmp_path / "vilextime/short_answers.jsonl").open(encoding="utf-8")]
    assert (s[0]["short_answer"], s[0]["model_short_answer"], s[1]["short_answer"]) == ("Phạt 7 triệu", "7 triệu", "9 triệu")


def test_label_ui_timeline_reads_the_version_chain_and_ends_it_at_the_repeal():
    import label_ui

    chain = ("[1] 2017-11-01 → 2020-03-22 · bản gốc   ◀ đáp án của dòng này\n"
             "[2] 2020-03-22 → nay · sửa bởi 17/2020/NĐ-CP (hiệu lực 2020-03-22)\n"
             "    chỗ đổi: bỏ «Chi»")
    t = label_ui.timeline({"cac_phien_ban": chain, "answer": "x"})
    assert [(v["n"], v["from"], v["to"], v["by"], v["mine"]) for v in t["versions"]] == [
        (1, "2017-11-01", "2020-03-22", "bản gốc", True),
        (2, "2020-03-22", None, "sửa bởi 17/2020/NĐ-CP (hiệu lực 2020-03-22)", False)]
    assert t["repeal"] is None
    # T6b: the provision's own window is open; its ancestor's repeal ends it, and the row after it has no version
    t = label_ui.timeline({"cac_phien_ban": "[1] 2012-06-01 → nay · bản gốc\nHết hiệu lực từ 2018-10-08 "
                                            "(điều khoản cấp trên bị bãi bỏ) bởi 140/2018/NĐ-CP",
                           "answer": "(không có đáp án: provision_repealed)"})
    assert t["versions"] == [{"n": 1, "from": "2012-06-01", "to": "2018-10-08", "by": "bản gốc", "mine": False}]
    assert t["repeal"] == {"on": "2018-10-08", "mode": "điều khoản cấp trên bị bãi bỏ", "by": "140/2018/NĐ-CP"}
    # T1: the sheet has no chain, so the row's own window is the one version
    t = label_ui.timeline({"cac_phien_ban": "", "answer": "x", "valid_from": "2013-07-01", "valid_to": ""})
    assert t["versions"] == [{"n": 1, "from": "2013-07-01", "to": None, "by": "bản duy nhất", "mine": True}]
    # the release's version id, read from the sheet alone
    assert label_ui.version_id({"lineage_id": "provision:u#k2#a", "cac_phien_ban": chain, "answer": "x"}) == "version:u#k2#a:1"
    assert label_ui.version_id({"lineage_id": "provision:u", "cac_phien_ban": "", "answer": "x"}) == "version:u:1"
    assert label_ui.version_id({"lineage_id": "provision:u", "cac_phien_ban": "[1] 2012-06-01 → nay · bản gốc",
                                "answer": "(không có đáp án: provision_repealed)"}) is None


def test_label_ui_compare_pairs_consecutive_answers_and_marks_both_sides():
    import label_ui

    def chain(mine):
        return "\n".join(f"[{n}] {a} → {b}{'   ◀ đáp án của dòng này' if n == mine else ''}"
                         for n, a, b in ((1, "2019-01-01", "2020-01-01"), (2, "2020-01-01", "nay")))
    rows = [{"lineage_id": "p", "answer": "Phạt 7 triệu đồng.", "cac_phien_ban": chain(1)},
            {"lineage_id": "p", "answer": "Phạt 9 triệu đồng.", "cac_phien_ban": chain(2)},
            {"lineage_id": "p", "answer": "Phạt 9 triệu đồng.", "cac_phien_ban": chain(2)},    # its mid-interval twin
            {"lineage_id": "q", "answer": "Một bản.", "cac_phien_ban": "", "valid_from": "2019-01-01"}]
    got = label_ui.compare(rows)
    assert list(got) == ["p"]                                   # one answer only: nothing to compare
    (pair,) = got["p"]["pairs"]
    old, new = (v["text"] for v in got["p"]["versions"])
    assert (pair["a"], pair["b"]) == (1, 2)
    assert [old[s:e] for s, e in pair["sa"]] == ["7"] and [new[s:e] for s, e in pair["sb"]] == ["9"]

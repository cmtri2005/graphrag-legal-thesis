"""Independent gold: the draw is stratified, blind to what people saw, and split 75/75 with 30 shared."""
import collections

import pytest

from independent_gold import STRATA, draw, reveal, stratum

FACETS = {"T1": ("T1", None, None), "T2-regular": ("T2", False, None), "T2-near": ("T2", True, None),
          "T3": ("T3", None, None), "T6-direct": ("T6", None, "direct"), "T6-ancestor": ("T6", None, "ancestor")}


def _rows():
    rows = []
    for s, (group, near, mode) in FACETS.items():
        for n in range(25):
            for day in ("2020-01-01", "2022-01-01"):
                rows.append({"id": f"{s}:{n}:{day}", "group": group, "near_identical": near, "repeal_mode": mode,
                             "lineage_id": f"{s}:{n}", "split": "test", "as_of_strategy": "giua_khoang", "as_of": day})
    return rows + [{**rows[0], "id": "dev", "split": "dev"}]


def test_the_draw_is_stratified_unseen_and_split_75_75_with_30_shared():
    seen = {f"{s}:{n}" for s in STRATA for n in range(5)}
    items = draw(_rows(), seen)
    assert collections.Counter(stratum(i["row"]) for i in items) == {s: 20 for s in STRATA}
    lineages = [i["row"]["lineage_id"] for i in items]
    assert len(set(lineages)) == 120 and not seen & set(lineages)
    who = collections.Counter(i["who"] for i in items)
    assert who == {"AB": 30, "A": 45, "B": 45}
    assert all(sum(i["who"] == "AB" and stratum(i["row"]) == s for i in items) == 5 for s in STRATA)
    assert [i["row"]["id"] for i in draw(_rows(), seen)] == [i["row"]["id"] for i in items]   # seeded


def test_too_few_unseen_lineages_stop_the_draw():
    with pytest.raises(SystemExit, match="T1"):
        draw(_rows(), {f"T1:{n}" for n in range(10)})


def test_reveal_waits_for_every_blind_answer(tmp_path):
    (tmp_path / "review").mkdir()
    (tmp_path / "review/gold_A.tsv").write_text(
        "khoa\tcau_hoi\tngay_hoi\tvan_ban\tdieu_khoan\tlink\tvan_ban_tac_dong_theo_cong\thieu_luc\tnoi_dung\t"
        "van_ban_ap_dung\tphut\tghi_chu\nG001\tq\t2020-01-01\tv\tĐiều 1\tu\t-\tcòn\t\t\t\t\n", encoding="utf-8")
    with pytest.raises(SystemExit, match="chưa xong"):
        reveal(tmp_path, "A")          # "còn" without the wording in force


def test_the_workbook_round_trip_loads_answers_and_refuses_an_edited_data_cell(tmp_path):
    from openpyxl import load_workbook

    from independent_gold import SHEET, export, ingest, read

    (tmp_path / "review").mkdir()
    head = "\t".join(SHEET)
    data_cells = ["G001", "q", "2020-01-01", "v", "Điều 1", "https://vbpl.vn/x", "-"]
    (tmp_path / "review/gold_A.tsv").write_text(f"{head}\n" + "\t".join(data_cells + [""] * 5) + "\n", encoding="utf-8")
    export(tmp_path, "A")
    book = tmp_path / "review/gold_A.xlsx"
    wb = load_workbook(book)
    ws = wb["Gán nhãn"]
    ws.cell(2, SHEET.index("hieu_luc") + 1, "còn")
    ws.cell(2, SHEET.index("noi_dung") + 1, "1. Nội dung đang có hiệu lực")
    wb.save(book)
    ingest(tmp_path, "A")
    assert read(tmp_path / "review/gold_A.tsv")[0]["noi_dung"] == "1. Nội dung đang có hiệu lực"
    ws.cell(2, SHEET.index("cau_hoi") + 1, "câu hỏi đã sửa")
    wb.save(book)
    with pytest.raises(SystemExit, match="cau_hoi"):
        ingest(tmp_path, "A")

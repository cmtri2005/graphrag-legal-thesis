#!/usr/bin/env python3
"""ViLexTime D4 — hoà giải các dòng A và B lệch nhau (W2-03, hướng dẫn §5 bước 4).

    export   -> data/review/hoa_giai_D4.xlsx: mỗi dòng lệch một hàng, nhãn và ghi
               chú của A và B đặt cạnh nhau; chỉ ba cột `chot_*` mở để điền.
    import   đọc xlsx, kiểm như nạp nhãn thường, rồi ghi
               data/review/hoa_giai_D4.tsv (id, verdict, reasons, note đã chốt) và
               cột `adjudicated` của crosscheck_A.tsv (= verdict đã chốt).
               Không ghi gì nếu còn một lỗi hoặc còn dòng chưa chốt.

Nhãn thô `verdict` của A và B không bị đụng tới: chúng được phát hành nguyên vẹn
(ADR 0003 §9), và AC1 đã đo trên chúng.

Usage:
    python scripts/review/adjudicate.py export
    python scripts/review/adjudicate.py import --by "human:A+B"

`--by` ghi ai chốt vào cột `adjudicator` (người, hoặc model kèm người xác nhận):
biên bản D5 phải nói đúng điều đó.
"""
from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from crosscheck import REASONS, VERDICTS, read  # noqa: E402
from crosscheck_xlsx import FONT, sheet_reasons  # noqa: E402

SHEET = "Hoà giải"
CHOT = ("chot_verdict", "chot_reasons", "chot_note")
COLUMNS = ["stt", "id", "group", "nguyen_nhan", "question", "as_of", "lead_in", "answer", "citation",
           "source_url", "A_verdict", "A_reasons", "A_note", "A_version",
           "B_verdict", "B_reasons", "B_note", "B_version", *CHOT]
WIDTHS = {"stt": 5, "id": 30, "group": 7, "nguyen_nhan": 20, "question": 55, "as_of": 11, "lead_in": 30,
          "answer": 55, "citation": 18, "source_url": 24, "A_verdict": 11, "A_reasons": 18, "A_note": 40,
          "A_version": 8, "B_verdict": 11, "B_reasons": 18, "B_note": 50, "B_version": 8,
          "chot_verdict": 13, "chot_reasons": 22, "chot_note": 40}
WRAP = {"question", "lead_in", "answer", "A_note", "B_note", "chot_note", "A_reasons", "B_reasons"}


def cause(a: dict, b: dict) -> str:
    if a["guideline_version"] != b["guideline_version"]:
        return "khác phiên bản HD"
    if a["group"] == "T6" and "gold_wrong" in (a["reasons"] + b["reasons"]):
        return "T6 chuyển chỗ?"
    return "hiểu khác"


def disagreements(data: Path) -> list[dict]:
    a_rows = read(data / "review/crosscheck_A.tsv")
    b = {r["id"]: r for r in read(data / "review/crosscheck_B.tsv")}
    out = []
    for a in a_rows:
        other = b[a["id"]]
        if a["verdict"] and other["verdict"] and a["verdict"] != other["verdict"]:
            row = {k: a[k] for k in ("id", "group", "question", "as_of", "lead_in", "answer",
                                     "citation", "source_url", "lineage_id")}
            row["nguyen_nhan"] = cause(a, other)
            for who, r in (("A", a), ("B", other)):
                for k in ("verdict", "reasons", "note"):
                    row[f"{who}_{k}"] = r[k]
                row[f"{who}_version"] = r["guideline_version"]
            out.append(row)
    order = ("T6 chuyển chỗ?", "hiểu khác", "khác phiên bản HD")
    out.sort(key=lambda r: (order.index(r["nguyen_nhan"]), r["group"], r["lineage_id"], r["as_of"]))
    for i, r in enumerate(out, 1):
        r["stt"] = i
    return out


GUIDE = [
    ("Hoà giải kiểm chéo D4", 15, True),
    ("Mỗi hàng là một dòng A và B gán khác nhau. Hai người ngồi cùng nhau, đọc lại văn bản gốc, chốt một nhãn.", 10, False),
    ("", 10, False),
    ("Cách làm mỗi hàng", 12, True),
    ("1. Đọc câu hỏi, mốc as_of, đáp án. Đọc lý do và ghi chú của cả A và B.", 10, False),
    ("2. Mở source_url, kiểm chỗ người kia đã chỉ ra (thường là chỗ mới của nội dung ở B_note).", 10, False),
    ("3. Điền chot_verdict (dropdown), chot_reasons nếu khong_dat, chot_note khi gold_wrong/other.", 10, False),
    ("   Theo hướng dẫn 2.0 hiện hành, kể cả với dòng trước đây gán theo 1.2.", 10, False),
    ("", 10, False),
    ("Cột nguyen_nhan (máy đoán, chỉ để sắp xếp)", 12, True),
    ("T6 chuyển chỗ?      dòng T6 có một người ghi gold_wrong: kiểm bảy yếu tố chuyển chỗ (hướng dẫn §3).", 10, False),
    ("hiểu khác           cùng phiên bản hướng dẫn, hai người hiểu khác.", 10, False),
    ("khác phiên bản HD   một người gán theo 1.2, người kia theo 2.0.", 10, False),
    ("", 10, False),
    ("Lưu ý", 12, True),
    ("· Không sửa cột nào ngoài ba cột chot_*. Nhãn gốc của A và B giữ nguyên, không đổi.", 10, False),
    ("· Dòng khong_dat sau hoà giải sẽ bị loại hoặc viết lại câu hỏi; không ai sửa tay đáp án.", 10, False),
    ("· Xong thì gửi lại file này để chạy: python scripts/review/adjudicate.py import", 10, False),
]


def export(data: Path) -> None:
    out = data / "review/hoa_giai_D4.xlsx"
    if out.exists():
        raise SystemExit(f"{out} đã có; không ghi đè kết quả hoà giải")
    rows = disagreements(data)
    wb = Workbook()
    ws = wb.active
    ws.title = "Đọc trước"
    for i, (text, size, bold) in enumerate(GUIDE, 1):
        ws.cell(i, 1, text).font = Font(name=FONT, size=size, bold=bold)
    ws.column_dimensions["A"].width = 112
    ws.sheet_view.showGridLines = False
    sheet_reasons(wb)

    ws = wb.create_sheet(SHEET)
    fills = {"A": "1F4E79", "B": "7F3F00", "c": "2F6B3A"}
    for col, name in enumerate(COLUMNS, 1):
        c = ws.cell(1, col, name)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=fills.get(name[0], "44546A"))
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = WIDTHS[name]
    editable = PatternFill("solid", fgColor="EAF3EC")
    for r, row in enumerate(rows, 2):
        for col, name in enumerate(COLUMNS, 1):
            c = ws.cell(r, col, row.get(name, ""))
            c.font = Font(name=FONT, size=10)
            c.alignment = Alignment(vertical="top", wrap_text=name in WRAP)
            c.protection = Protection(locked=name not in CHOT)
            if name in CHOT:
                c.fill = editable
            if name == "source_url" and row.get(name):
                c.hyperlink = row[name]
                c.font = Font(name=FONT, size=10, color="0563C1", underline="single")
    dv = DataValidation(type="list", formula1='"' + ",".join(sorted(VERDICTS)) + '"',
                        allow_blank=True, showErrorMessage=True)
    ws.add_data_validation(dv)
    letter = get_column_letter(COLUMNS.index("chot_verdict") + 1)
    dv.add(f"{letter}2:{letter}{len(rows) + 1}")
    ws.freeze_panes = ws.cell(2, COLUMNS.index("question") + 1).coordinate
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"
    ws.protection.sheet = True
    ws.protection.enable()
    wb.save(out)
    counts = {k: sum(r["nguyen_nhan"] == k for r in rows) for k in ("T6 chuyển chỗ?", "hiểu khác", "khác phiên bản HD")}
    print(f"{out}: {len(rows)} dòng lệch {counts}")


def check(rows: dict[str, dict], expected: set[str]) -> list[str]:
    errors = []
    if set(rows) != expected:
        errors.append(f"xlsx có {len(set(rows) - expected)} id lạ, thiếu {len(expected - set(rows))} id")
    for rid, cells in rows.items():
        verdict, note = cells["chot_verdict"], cells["chot_note"]
        codes = [c.strip() for c in cells["chot_reasons"].split(",") if c.strip()]
        if verdict not in VERDICTS:
            errors.append(f"{rid}: chot_verdict `{verdict}` trống hoặc ngoài từ vựng")
        errors += [f"{rid}: chot_reasons `{c}` ngoài từ vựng" for c in codes if c not in REASONS]
        if verdict == "khong_dat" and not codes:
            errors.append(f"{rid}: khong_dat nhưng thiếu chot_reasons")
        if verdict == "dat" and codes:
            errors.append(f"{rid}: dat nhưng vẫn có chot_reasons")
        if {"gold_wrong", "other"} & set(codes) and not note:
            errors.append(f"{rid}: gold_wrong/other bắt buộc có chot_note")
    return errors


def ingest(data: Path, by: str) -> None:
    ws = load_workbook(data / "review/hoa_giai_D4.xlsx", data_only=True)[SHEET]
    head = [c.value for c in ws[1]]
    rows = {}
    for values in ws.iter_rows(min_row=2, values_only=True):
        cells = {k: str(v or "").strip() for k, v in zip(head, values)}
        if cells.get("id"):
            rows[cells["id"]] = cells
    errors = check(rows, {r["id"] for r in disagreements(data)})
    if errors:
        print(f"{len(errors)} lỗi — KHÔNG ghi gì:", file=sys.stderr)
        for e in errors[:40]:
            print(f"  {e}", file=sys.stderr)
        raise SystemExit(1)
    with (data / "review/hoa_giai_D4.tsv").open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter="\t", lineterminator="\n")
        w.writerow(["id", "verdict", "reasons", "note", "adjudicator"])
        for rid, c in rows.items():
            w.writerow([rid, c["chot_verdict"], c["chot_reasons"], c["chot_note"], by])
    tsv = data / "review/crosscheck_A.tsv"
    a_rows = read(tsv)
    for r in a_rows:
        r["adjudicated"] = rows[r["id"]]["chot_verdict"] if r["id"] in rows else r.get("adjudicated", "")
    with tsv.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(a_rows[0]), delimiter="\t", lineterminator="\n")
        w.writeheader()
        w.writerows(a_rows)
    final = [c["chot_verdict"] for c in rows.values()]
    print(f"chốt {len(rows)} dòng: dat {final.count('dat')}, khong_dat {final.count('khong_dat')} "
          f"-> data/review/hoa_giai_D4.tsv, cột adjudicated của {tsv}")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("command", choices=("export", "import"))
    ap.add_argument("--data", type=Path, default=Path("data"))
    ap.add_argument("--by", help="bắt buộc với import: ai chốt")
    args = ap.parse_args()
    if args.command == "export":
        export(args.data)
    elif not args.by:
        ap.error("import cần --by")
    else:
        ingest(args.data, args.by)


if __name__ == "__main__":
    main()

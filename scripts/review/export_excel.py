#!/usr/bin/env python3
"""ViLexTime — the current release selection as an Excel workbook, for people.

A read-only view for the team: it writes nothing back and the repository stays
the system of record (ADR 0002). Re-run whenever `selected.jsonl` changes.

Tabs:
  Đọc trước   status of this preview, what each column means
  Dataset     one row per QA record (ADR 0003), dates and ids kept as text
  Thống kê    per-group counts as formulas over the Dataset tab

Usage:
    python scripts/review/export_excel.py
    python scripts/review/export_excel.py --out somewhere.xlsx
"""
from __future__ import annotations

import argparse
import json
import subprocess
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT = "Arial"
QUOTA = {"T1": 250, "T2": 400, "T3": 200, "T4": 150, "T5": 50, "T6": 100}
GROUPS = {
    "T1": "Một phiên bản (điều khoản chưa từng bị sửa)",
    "T2": "Hai phiên bản khác rõ (cặp tương phản trước/sau sửa đổi)",
    "T3": "Chuỗi ≥ 3 phiên bản",
    "T4": "Hai phiên bản gần giống (bẫy phiên bản cũ)",
    "T5": "Hiệu lực trở về trước (chờ luật sư)",
    "T6": "Hết hiệu lực một phần ở cấp khoản, điểm",
}
# (header, width, wrap, value-from-record, meaning)
COLUMNS = [
    ("id", 30, False, lambda r: r["id"], "Khoá dòng: nhóm, điều khoản, mốc hỏi"),
    ("group", 7, False, lambda r: r["group"], "Nhóm T1–T6"),
    ("question", 50, True, lambda r: r["question"], "Câu hỏi (LLM viết, đã qua kiểm tự động, chưa qua người)"),
    ("as_of", 11, False, lambda r: r["as_of"], "Mốc thời gian câu hỏi được hỏi (YYYY-MM-DD)"),
    ("answer", 70, True, lambda r: r["answer"] if r["answer"] is not None else "",
     "Đáp án chuẩn: nguyên văn điều khoản có hiệu lực tại as_of; rỗng khi không có đáp án"),
    ("answerable", 11, False, lambda r: "có" if r["answerable"] else "không", "Có đáp án tại as_of không"),
    ("no_answer_reason", 18, False, lambda r: r["no_answer_reason"] or "",
     "Lý do không có đáp án (provision_repealed = điều khoản đã bị bãi bỏ)"),
    ("law_id", 20, False, lambda r: r["target"]["law_id"], "Số hiệu văn bản chứa điều khoản"),
    ("citation", 24, False, lambda r: r["target"]["citation"], "Điều/Khoản/Điểm"),
    ("provision_path", 30, True, lambda r: r["target"]["provision_path"], "Vị trí trong cây văn bản"),
    ("lead_in", 30, True, lambda r: r["target"].get("lead_in") or "", "Tiêu đề/câu dẫn của nút cha"),
    ("document_title", 40, True, lambda r: " ".join((r["target"]["document_title"] or "").split()), "Tên văn bản"),
    ("version_id", 22, False, lambda r: r["version_id"] or "", "Phiên bản đáp án"),
    ("valid_from", 11, False, lambda r: r["valid_from"] or "", "Phiên bản có hiệu lực từ (nửa mở [from, to))"),
    ("valid_to", 11, False, lambda r: r["valid_to"] or "", "Phiên bản hết hiệu lực từ; rỗng = còn mở"),
    ("num_versions", 9, False, lambda r: r["num_versions"], "Số phiên bản trong chuỗi"),
    ("amendments", 40, True,
     lambda r: "; ".join(f"{t['on']}: {t['actor']['law_id']} (v{t['from_version']}→v{t['to_version']})"
                         for t in r["derivation"]["transitions"]),
     "Các lần sửa đổi: ngày, văn bản sửa đổi, phiên bản"),
    ("retroactive", 34, True,
     lambda r: (f"{r['retroactive']['actor_law_id']}: ký {r['retroactive']['issued_on']}, "
                f"hiệu lực {r['retroactive']['effective_from']}") if r.get("retroactive") else "",
     "Chỉ T5: văn bản hồi tố, ngày ký, ngày hiệu lực"),
    ("source_url", 16, False, lambda r: r["source_url"], "Link văn bản trên vbpl.vn"),
    ("review_status", 12, False, lambda r: r["review"]["status"], "unchecked = chưa ai kiểm"),
    ("split", 8, False, lambda r: r["split"] or "", "dev/test — chưa chia"),
]
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", start_color="1F4E78")
BODY = Font(name=FONT, size=10)
TOP = Alignment(vertical="top")
WRAP = Alignment(vertical="top", wrap_text=True)


def header(ws, row: int, values: list[str]) -> None:
    for col, value in enumerate(values, 1):
        c = ws.cell(row, col, value)
        c.font, c.fill, c.alignment = HEAD, HEAD_FILL, WRAP


def dataset(wb: Workbook, rows: list[dict]) -> None:
    ws = wb.create_sheet("Dataset")
    header(ws, 1, [name for name, *_ in COLUMNS])
    url_col = [name for name, *_ in COLUMNS].index("source_url") + 1
    for i, r in enumerate(rows, 2):
        for col, (_, _, wrap, get, _) in enumerate(COLUMNS, 1):
            c = ws.cell(i, col, get(r))
            c.font, c.alignment = BODY, WRAP if wrap else TOP
            if isinstance(c.value, str):
                c.number_format = "@"          # keep 2025-06-30 and 25.000.000 as written
        link = ws.cell(i, url_col)
        link.hyperlink, link.font = link.value, Font(name=FONT, size=10, color="0563C1", underline="single")
    for col, (_, width, *_) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"


def statistics(wb: Workbook, n: int) -> None:
    ws = wb.create_sheet("Thống kê")
    header(ws, 1, ["Nhóm", "Mô tả", "Quota đề cương", "Số dòng", "Có đáp án", "Không có đáp án", "Chênh quota"])
    last = n + 1
    for i, (g, desc) in enumerate(GROUPS.items(), 2):
        ws.cell(i, 1, g); ws.cell(i, 2, desc); ws.cell(i, 3, QUOTA[g])
        ws.cell(i, 4, f"=COUNTIF(Dataset!$B$2:$B${last},A{i})")
        ws.cell(i, 5, f'=COUNTIFS(Dataset!$B$2:$B${last},A{i},Dataset!$F$2:$F${last},"có")')
        ws.cell(i, 6, f"=D{i}-E{i}")
        ws.cell(i, 7, f"=D{i}-C{i}")
    t = len(GROUPS) + 2
    ws.cell(t, 1, "Tổng")
    for col in "CDEFG":
        ws[f"{col}{t}"] = f"=SUM({col}2:{col}{t - 1})"
    for row in ws.iter_rows(min_row=2, max_row=t):
        for c in row:
            c.font = Font(name=FONT, size=10, bold=c.row == t)
    ws.cell(t + 2, 1, "T6 nhận phần T5 thiếu (T5 chờ luật sư) để giữ tổng 1.150 — quyết định trong plan phase-4.").font = BODY
    for col, width in zip("ABCDEFG", (8, 52, 15, 10, 11, 16, 12)):
        ws.column_dimensions[col].width = width


def readme(wb: Workbook, n: int, commit: str) -> None:
    ws = wb.active
    ws.title = "Đọc trước"
    lines = [
        ("ViLexTime — bản xem trước", True),
        (f"{n} dòng QA · schema 1.0.0 (ADR 0003) · dựng từ commit {commit}", False),
        ("", False),
        ("TRẠNG THÁI: CHƯA PHÁT HÀNH", True),
        ("• Câu hỏi mới qua kiểm tự động; chưa qua người kiểm (review_status = unchecked).", False),
        ("• Kiểm chéo Cohen κ trên 300 dòng: chưa gán nhãn (data/review/crosscheck_A.tsv, _B.tsv).", False),
        ("• T5: chờ luật sư xét theo Điều 152 Luật BHVBQPPL.", False),
        ("• Chưa chia dev/test (cột split rỗng).", False),
        ("• Đáp án chuẩn lấy từ dữ liệu phiên bản, không do LLM viết.", False),
        ("", False),
        ("KHÔNG sửa dữ liệu trong file này. Repository là bản gốc; file này sinh lại bằng", False),
        ("python scripts/review/export_excel.py", False),
        ("", False),
        ("CÁC CỘT (tab Dataset)", True),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        ws.cell(i, 1, text).font = Font(name=FONT, size=12 if i == 1 else 10, bold=bold)
    start = len(lines) + 1
    header(ws, start, ["Cột", "Ý nghĩa"])
    for i, (name, _, _, _, meaning) in enumerate(COLUMNS, start + 1):
        ws.cell(i, 1, name).font = Font(name=FONT, size=10, bold=True)
        ws.cell(i, 2, meaning).font = BODY
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 90


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--data", type=Path, default=Path("data"))
    ap.add_argument("--out", type=Path, default=Path("data/derived/vilextime/vilextime_preview.xlsx"))
    args = ap.parse_args()
    rows = [json.loads(line) for line in (args.data / "derived/vilextime/selected.jsonl").open(encoding="utf-8")]
    rows.sort(key=lambda r: (r["group"], r["lineage_id"], r["as_of"]))
    commit = subprocess.run(["git", "rev-parse", "--short", "HEAD"], capture_output=True, text=True).stdout.strip()
    wb = Workbook()
    readme(wb, len(rows), commit or "?")
    dataset(wb, rows)
    statistics(wb, len(rows))
    wb.save(args.out)
    print(f"{len(rows)} dòng -> {args.out}")


if __name__ == "__main__":
    main()

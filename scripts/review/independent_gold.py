#!/usr/bin/env python3
"""ViLexTime independent gold (MASTER_RESEARCH_GUIDE §5.3.5, decided 11/10: B6) — sample and worksheets.

The gold of every row is derived by the KG, so checking the KG against itself proves
nothing. Here a person reads the raw documents and decides, without seeing the KG's
answer, whether the provision is in force at the date and what it says. One piece of
work answers three questions: how often the KG is right (M2, threshold 114/120), how
often the gold is wrong, and whether the main comparison holds on a gold the KG did
not write (ADR 0005 §7).

The sample: 120 rows of `test` at `giua_khoang`, 20 per stratum (T1, T2 regular, T2
near-identical, T3, T6 direct repeal, T6 ancestor repeal), one row per lineage, and
none from a lineage a person has already seen with the KG's answer: D4, the Excel
preview (`human:` labels) and the audit queues. 30 rows, 5 per stratum, go to both
annotators; the other 90 split evenly, so each labels 75.

Labelling is blind, then revealed:

    build   writes data/review/gold_sample.tsv (key -> row; not for the annotators) and
            gold_A.tsv, gold_B.tsv: an opaque key, the question, the date, where to
            look (document, citation, portal link) and the documents the portal's Lược
            đồ tab lists as acting on that document. Never the KG's answer, its dates,
            the row id or the group, which would say "this one was repealed". The
            annotator fills `hieu_luc` (còn | không | không rõ), `noi_dung` (the wording
            in force, pasted from the source), `van_ban_ap_dung` (the document whose
            wording applies), `phut` and `ghi_chu`.
    export  gold_<X>.xlsx for the annotator, as D4 does (crosscheck_xlsx.py): a sheet
            to read first, and the rows with every data column locked, `hieu_luc` a
            closed dropdown and only the five answer columns open.
    import  reads gold_<X>.xlsx back into gold_<X>.tsv, all or nothing: no answer is
            loaded when a data cell was changed or a value is outside the vocabulary.
    reveal  once every row of gold_<X>.tsv is filled: gold_<X>.reveal.tsv, with the
            locked answers, the sha256 of the file they came from, the KG's answer
            beside them, and empty `dong_y`, `ly_do`, `ghi_chu_mo` cells.

No command overwrites a file that already carries a label.

Usage:
    python scripts/review/independent_gold.py build
    python scripts/review/independent_gold.py export A
    python scripts/review/independent_gold.py import A
    python scripts/review/independent_gold.py reveal A
"""
from __future__ import annotations

import argparse
import collections
import csv
import hashlib
import json
import random
from pathlib import Path

SEED = 20261011
PER_STRATUM, SHARED = 20, 5
STRATA = ("T1", "T2-regular", "T2-near", "T3", "T6-direct", "T6-ancestor")
URL = "https://vbpl.vn/TW/Pages/vbpq-toanvan.aspx?ItemID={}"
# Lược đồ codes as seen from the acting document (data/reference_type_map.json). Detailed
# guidance (9) and citations change no validity, so they are left out.
ACTS = {1: "bãi bỏ", 5: "đình chỉ thi hành", 6: "đính chính", 7: "hợp nhất", 10: "sửa đổi, bổ sung",
        11: "tạm ngưng hiệu lực", 12: "thay thế"}
FORCE = {"còn", "không", "không rõ"}
REASONS = {"kg_missed_amendment", "kg_wrong_dates", "kg_wrong_text", "kg_wrong_repeal", "annotator_error",
           "transitional_rule", "ambiguous", "other"}
SHEET = ["khoa", "cau_hoi", "ngay_hoi", "van_ban", "dieu_khoan", "link", "van_ban_tac_dong_theo_cong",
         "hieu_luc", "noi_dung", "van_ban_ap_dung", "phut", "ghi_chu"]
FILLED = SHEET[7:]
REVEAL = SHEET[:6] + FILLED[:3] + ["stage1_sha256", "kg_hieu_luc", "kg_noi_dung", "kg_hieu_luc_tu",
                                   "kg_hieu_luc_den", "dong_y", "ly_do", "ghi_chu_mo"]


def stratum(r: dict) -> str:
    if r["group"] == "T2":
        return "T2-near" if r.get("near_identical") else "T2-regular"
    if r["group"] == "T6":
        return f"T6-{r.get('repeal_mode') or 'direct'}"
    return r["group"]


def seen_lineages(review: Path, lineage_of: dict[str, str]) -> set[str]:
    """Lineages a person has already seen together with the KG's answer."""
    seen = set()
    def lineage(row_id: str) -> str:   # an id from before 07/10 may be gone from the release
        return lineage_of.get(row_id) or row_id.split(":", 2)[-1].rsplit(":", 1)[0]
    for name in ("crosscheck_A.tsv", "crosscheck_B.tsv", "audit_queue.tsv", "audit_queue_relocation.tsv",
                 "audit_truoc_hieu_luc.tsv", "labels.tsv"):
        path = review / name
        if not path.exists():
            continue
        for r in csv.DictReader(path.open(encoding="utf-8"), delimiter="\t"):
            if name == "labels.tsv" and not r.get("labeler", "").startswith("human"):
                continue   # a model label was never shown to a person
            seen.add(r.get("lineage_id") or lineage(r["id"]))
    return seen


def draw(rows: list[dict], seen: set[str]) -> list[dict]:
    """[{key, row, who}]: 20 lineages per stratum, one row each; 5 of them to both annotators."""
    pool = collections.defaultdict(lambda: collections.defaultdict(list))
    for r in rows:
        if (r.get("split") == "test" and r["as_of_strategy"] == "giua_khoang" and stratum(r) in STRATA
                and r["lineage_id"] not in seen):
            pool[stratum(r)][r["lineage_id"]].append(r)
    picked, alone = [], []
    for s in STRATA:
        lineages = sorted(pool[s])
        if len(lineages) < PER_STRATUM:
            raise SystemExit(f"{s}: chỉ còn {len(lineages)} lineage chưa ai xem, cần {PER_STRATUM}")
        random.Random(f"{SEED}:{s}").shuffle(lineages)
        for i, lineage in enumerate(lineages[:PER_STRATUM]):
            row = random.Random(f"{SEED}:{lineage}").choice(sorted(pool[s][lineage], key=lambda r: r["as_of"]))
            (picked if i < SHARED else alone).append({"row": row, "who": "AB"})
    for i, item in enumerate(alone):          # strata interleave, so each annotator gets every stratum
        item["who"] = "AB"[i % 2]
    out = picked + alone
    random.Random(SEED).shuffle(out)
    for i, item in enumerate(out, 1):
        item["key"] = f"G{i:03d}"
    return out


def actors(data: Path, document_id: str) -> str:
    """What the portal's Lược đồ tab lists as acting on the document: a pointer, not an answer."""
    path = data / "diagrams" / f"{document_id.removeprefix('document:')}.json"
    if not path.exists():
        return "(không có dữ liệu Lược đồ — tra trực tiếp trên cổng)"
    diagram = json.loads(path.read_text(encoding="utf-8"))
    found = [f"{ACTS[int(code)]}: {' '.join(item['name'].split())[:140]} ({URL.format(item['id'])})"
             for code, items in (diagram.get("documentNamesBySource") or {}).items() if int(code) in ACTS
             for item in items or []]
    return " | ".join(found) or "(cổng không liệt kê văn bản tác động nào)"


def read(path: Path) -> list[dict]:
    return list(csv.DictReader(path.open(encoding="utf-8", newline=""), delimiter="\t")) if path.exists() else []


def write(path: Path, columns: list[str], rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, columns, delimiter="\t", lineterminator="\n")
        w.writeheader()
        w.writerows(rows)


def refuse_if_labelled(path: Path, cells: list[str]) -> None:
    if any(r.get(c) for r in read(path) for c in cells):
        raise SystemExit(f"{path} đã có nhãn; không ghi đè")


def build(data: Path) -> None:
    review = data / "review"
    for who in "AB":
        refuse_if_labelled(review / f"gold_{who}.tsv", FILLED)
    rows = [json.loads(line) for line in (data / "derived/vilextime/release.jsonl").open(encoding="utf-8")]
    items = draw(rows, seen_lineages(review, {r["id"]: r["lineage_id"] for r in rows}))
    write(review / "gold_sample.tsv", ["khoa", "id", "lineage_id", "stratum", "annotators"],
          [{"khoa": i["key"], "id": i["row"]["id"], "lineage_id": i["row"]["lineage_id"],
            "stratum": stratum(i["row"]), "annotators": i["who"]} for i in items])
    for who in "AB":
        mine = [i for i in items if who in i["who"]]
        write(review / f"gold_{who}.tsv", SHEET, [{
            "khoa": i["key"], "cau_hoi": i["row"]["question"], "ngay_hoi": i["row"]["as_of"],
            "van_ban": " ".join(i["row"]["target"]["document_title"].split()),
            "dieu_khoan": i["row"]["target"]["citation"], "link": i["row"]["source_url"],
            "van_ban_tac_dong_theo_cong": actors(data, i["row"]["target"]["document_id"])} for i in mine])
        print(f"gold_{who}.tsv: {len(mine)} dòng")
    print(f"{len(items)} dòng, {sum(i['who'] == 'AB' for i in items)} gán đôi -> {review}")


GUIDE = [
    "Gold độc lập — bước một: gán mù",
    "Mục đích: kiểm xem đồ thị tri thức (KG) có xác định đúng điều khoản có hiệu lực hay không, bằng cách tự đọc văn bản gốc.",
    "Không mở đáp án của hệ thống, release.jsonl hay data/review/gold_sample.tsv trước khi nộp bước này. Thấy đáp án trước là mất tính độc lập.",
    "",
    "Với mỗi dòng:",
    "1. Đọc `cau_hoi` và `ngay_hoi`. Mở `link` (văn bản gốc trên vbpl.vn), tìm đúng điều khoản ở cột `dieu_khoan`.",
    "2. Cột `van_ban_tac_dong_theo_cong` là các văn bản mà tab Lược đồ của cổng liệt kê là sửa đổi, thay thế, bãi bỏ… văn bản này.",
    "   Mở từng văn bản có hiệu lực trước ngày hỏi, xem nó có chạm điều khoản đang hỏi hoặc Điều, Chương chứa nó không.",
    "   Danh sách này có thể thiếu. Nghi thì tra thêm trên cổng.",
    "3. `hieu_luc`: còn = điều khoản đang có hiệu lực tại ngày hỏi; không = đã bị bãi bỏ, thay thế, hết hiệu lực, hoặc chưa có hiệu lực;",
    "   không rõ = không xác định được (ghi lý do ở `ghi_chu`).",
    "4. Nếu còn: dán vào `noi_dung` đúng lời văn đang có hiệu lực tại ngày hỏi (bản gốc hoặc bản đã sửa), ghi số hiệu văn bản có lời văn đó",
    "   vào `van_ban_ap_dung`.",
    "5. Điều khoản chuyển tiếp giữ quy định cũ cho trường hợp trong câu hỏi: vẫn điền theo quy định được áp dụng, và ghi Điều chuyển tiếp vào `ghi_chu`.",
    "6. `phut`: số phút đã dùng cho dòng này.",
    "",
    "Chỉ năm cột được mở để điền; các cột khác đã khóa. Xong thì báo nhóm chạy `import` rồi `reveal`.",
    "Sau `reveal`, câu trả lời bước một bị khóa bằng sha256 và không sửa nữa; bước hai mới cho xem đáp án của KG.",
]
WIDTHS = {"khoa": 7, "cau_hoi": 60, "ngay_hoi": 11, "van_ban": 40, "dieu_khoan": 18, "link": 28,
          "van_ban_tac_dong_theo_cong": 60, "hieu_luc": 10, "noi_dung": 60, "van_ban_ap_dung": 20, "phut": 7,
          "ghi_chu": 40}


def export(data: Path, who: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Protection
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    rows, out = read(data / f"review/gold_{who}.tsv"), data / f"review/gold_{who}.xlsx"
    if out.exists():
        have = {r["khoa"]: [r.get(c, "").strip() for c in FILLED] for r in rows}
        pending = [x["khoa"] for x in _from_xlsx(out)
                   if any(x.get(c) for c in FILLED) and [x.get(c, "") for c in FILLED] != have.get(x["khoa"])]
        if pending:
            raise SystemExit(f"{out} có {len(pending)} dòng chưa nạp về TSV; chạy import trước")
    wb = Workbook()
    guide = wb.active
    guide.title = "Đọc trước"
    for i, line in enumerate(GUIDE, 1):
        guide.cell(i, 1, line).font = Font(size=13 if i == 1 else 11, bold=i == 1)
    guide.column_dimensions["A"].width = 140
    ws = wb.create_sheet(LABEL_SHEET)
    for col, name in enumerate(SHEET, 1):
        c = ws.cell(1, col, name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F6B3A" if name in FILLED else "44546A")
        ws.column_dimensions[get_column_letter(col)].width = WIDTHS[name]
    for r, row in enumerate(rows, 2):
        for col, name in enumerate(SHEET, 1):
            c = ws.cell(r, col, row.get(name) or "")
            c.alignment = Alignment(vertical="top", wrap_text=name in ("cau_hoi", "van_ban_tac_dong_theo_cong",
                                                                        "noi_dung", "ghi_chu", "van_ban"))
            c.protection = Protection(locked=name not in FILLED)
            if name in FILLED:
                c.fill = PatternFill("solid", fgColor="EAF3EC")
            if name == "link" and row.get(name):
                c.hyperlink, c.font = row[name], Font(color="0563C1", underline="single")
    force = DataValidation(type="list", formula1='"' + ",".join(sorted(FORCE)) + '"', allow_blank=True,
                           showErrorMessage=True)
    ws.add_data_validation(force)
    letter = get_column_letter(SHEET.index("hieu_luc") + 1)
    force.add(f"{letter}2:{letter}{len(rows) + 1}")
    ws.freeze_panes = "C2"
    ws.protection.sheet = True       # a guard against slips, not security
    wb.save(out)
    print(f"{out}: {len(rows)} dòng")


LABEL_SHEET = "Gán nhãn"


def _from_xlsx(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    ws = load_workbook(path, data_only=True)[LABEL_SHEET]
    head = [c.value for c in ws[1]]
    return [{k: "" if v is None else str(v).strip() for k, v in zip(head, row)}
            for row in ws.iter_rows(min_row=2, values_only=True) if row[0]]


def ingest(data: Path, who: str) -> None:
    tsv = data / f"review/gold_{who}.tsv"
    rows, got = read(tsv), {r["khoa"]: r for r in _from_xlsx(data / f"review/gold_{who}.xlsx")}
    errors = [f"xlsx có {len(got)} dòng, tsv có {len(rows)}"] if len(got) != len(rows) else []
    for r in rows:
        x = got.get(r["khoa"])
        if x is None:
            errors.append(f"{r['khoa']}: thiếu trong xlsx")
            continue
        errors += [f"{r['khoa']}: cột `{c}` đã bị sửa" for c in SHEET if c not in FILLED and x.get(c, "") != r[c].strip()]
        if x["hieu_luc"] and x["hieu_luc"] not in FORCE:
            errors.append(f"{r['khoa']}: hieu_luc `{x['hieu_luc']}` ngoài {sorted(FORCE)}")
        if x["hieu_luc"] == "còn" and not x["noi_dung"]:
            errors.append(f"{r['khoa']}: còn hiệu lực nhưng thiếu noi_dung")
        if x["phut"] and not x["phut"].replace(".", "", 1).isdigit():
            errors.append(f"{r['khoa']}: phut `{x['phut']}` không phải số")
    if errors:
        raise SystemExit(f"{len(errors)} lỗi, không ghi gì vào {tsv}:\n  " + "\n  ".join(errors[:40]))
    write(tsv, SHEET, [{**r, **{c: got[r["khoa"]][c] for c in FILLED}} for r in rows])
    print(f"{tsv}: {sum(bool(got[r['khoa']]['hieu_luc']) for r in rows)}/{len(rows)} dòng đã có hieu_luc")


def reveal(data: Path, who: str) -> None:
    review = data / "review"
    sheet = review / f"gold_{who}.tsv"
    answers = read(sheet)
    missing = [r["khoa"] for r in answers if r["hieu_luc"] not in FORCE or (r["hieu_luc"] == "còn" and not r["noi_dung"].strip())]
    if not answers or missing:
        raise SystemExit(f"{sheet}: chưa xong, còn {len(missing)} dòng thiếu hiệu lực hoặc nội dung: {missing[:10]}")
    out = review / f"gold_{who}.reveal.tsv"
    refuse_if_labelled(out, ["dong_y", "ly_do", "ghi_chu_mo"])
    lock = hashlib.sha256(sheet.read_bytes()).hexdigest()
    row_of = {r["khoa"]: r["id"] for r in read(review / "gold_sample.tsv")}
    release = {r["id"]: r for r in map(json.loads, (data / "derived/vilextime/release.jsonl").open(encoding="utf-8"))}
    lines = []
    for a in answers:
        kg = release[row_of[a["khoa"]]]
        lines.append({**{c: a[c] for c in REVEAL if c in a}, "stage1_sha256": lock,
                      "kg_hieu_luc": "còn" if kg["answerable"] else "không",
                      "kg_noi_dung": kg["answer"] if kg["answerable"] else f"(không có đáp án: {kg['no_answer_reason']})",
                      "kg_hieu_luc_tu": kg["valid_from"] or "", "kg_hieu_luc_den": kg["valid_to"] or ""})
    write(out, REVEAL, lines)
    print(f"{out}: {len(lines)} dòng; bước một khóa ở sha256 {lock[:12]}…; `ly_do` thuộc {sorted(REASONS)}")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("command", choices=("build", "export", "import", "reveal"))
    ap.add_argument("annotator", nargs="?", choices=("A", "B"))
    ap.add_argument("--data", type=Path, default=Path("data"))
    args = ap.parse_args()
    if args.command != "build" and not args.annotator:
        ap.error(f"{args.command} needs A or B")
    if args.command == "build":
        build(args.data)
    else:
        {"export": export, "import": ingest, "reveal": reveal}[args.command](args.data, args.annotator)


if __name__ == "__main__":
    main()
